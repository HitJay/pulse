"""
RIC Monthly Report — August 2026 (2026-07-28 -> 2026-08-27)
Derived from build_ric_monthly_2026_07.py (fill_cell inline-run semantics, hyperlink auto-derivation intact).
Run: conda run -n ppt_env python scripts/ric-monthly/build_ric_monthly_2026_08.py
(from /das/user/QYJI/pulse repo root)
"""

import re

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

JIRA_BASE = "https://jira.novonordisk.com/browse/"
RIC_RE = re.compile(r"RIC-\d+")

# ── Colours ──
NN_BLUE      = RGBColor(0x00, 0x38, 0x65)
NN_BLUE_LITE = RGBColor(0xCC, 0xD9, 0xE8)
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GREY    = RGBColor(0x26, 0x26, 0x26)
GREY_BG      = RGBColor(0xF2, 0xF2, 0xF2)

# ── Font sizes (user-confirmed 2026-05-28, do not shrink) ──
SZ_HEADER    = 11
SZ_BODY      = 9.5
SZ_LINK      = 8.5
SZ_TITLE_BAR = 14
LINE_SPACING = 12

def set_cell_bg(cell, rgb):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for tag in ['a:solidFill', 'a:gradFill', 'a:noFill']:
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    solidFill = etree.SubElement(tcPr, qn('a:solidFill'))
    srgb = etree.SubElement(solidFill, qn('a:srgbClr'))
    srgb.set('val', f'{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}')

def add_run(para, text, bold=False, size_pt=SZ_BODY, color=DARK_GREY, italic=False, hyperlink=None):
    run = para.add_run()
    run.text = text
    run.font.bold = bold
    run.font.italic = italic
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.name = 'Calibri'
    if hyperlink:
        run.hyperlink.address = hyperlink
    return run

def _style_para(para):
    para.space_before = Pt(0.5)
    para.space_after  = Pt(0.5)
    para.line_spacing = Pt(LINE_SPACING)

def fill_cell(cell, lines, header=False, category=False, link=False):
    tf = cell.text_frame
    tf.word_wrap = True
    tf.margin_left   = Inches(0.04)
    tf.margin_right  = Inches(0.04)
    tf.margin_top    = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    c = WHITE if header else (NN_BLUE if category else DARK_GREY)
    sz = SZ_HEADER if header else (SZ_BODY if not link else SZ_LINK)

    para = tf.paragraphs[0]
    _style_para(para)
    para_used = False

    for (text, bold, italic) in lines:
        if link:
            if para_used:
                para = tf.add_paragraph(); _style_para(para)
            m = RIC_RE.search(text)
            href = JIRA_BASE + m.group(0) if m else None
            link_color = RGBColor(0x00, 0x56, 0x9E) if href else c
            add_run(para, text, bold=bold, italic=italic, size_pt=sz, color=link_color, hyperlink=href)
            para_used = True
            continue

        segments = text.split('\n')
        for si, seg in enumerate(segments):
            if si > 0:
                para = tf.add_paragraph(); _style_para(para)
                para_used = False
            if seg == '' and not para_used:
                continue
            add_run(para, seg, bold=bold, italic=italic, size_pt=sz, color=c)
            para_used = True

def vmerge_restart(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for el in tcPr.findall(qn('a:vMerge')):
        tcPr.remove(el)
    vm = etree.SubElement(tcPr, qn('a:vMerge'))
    vm.set('val', 'restart')

def vmerge_cont(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for el in tcPr.findall(qn('a:vMerge')):
        tcPr.remove(el)
    etree.SubElement(tcPr, qn('a:vMerge'))

# ── Table data ──
ROWS = [
    {
        "cat": "Capability", "cat_span": 3,
        "act": "Safety Landscape Tool — v1.4 MoA biology layer + operations hardening",
        "desc": [
            ("1. New ", False, False), ("Mechanism-of-Action biology layer", True, False),
            (" on OpenTargets v4: pathways, drug MoA, probes, subcellular location; action-type vocabulary completed from a live OT probe (BLOCKER / OPENER / POSITIVE MODULATOR...); direction-neutral terms stay silent.", False, False),
            ("\n2. Rendered as \"MoA Profile\" section in md/marp/pptx outputs; NCBN-305 gold-standard gate re-run: ", False, False),
            ("\u0394acc = 0", True, False),
            (" (scoring untouched).", False, False),
            ("\n3. ", False, False), ("Gold-standard direction-provenance fix", True, False),
            (" (HMGB1 request trigger) + snapshot-consistency audit.", False, False),
            ("\n4. Model alias migrated opus_latest \u2192 ", False, False), ("opus_5_azure", True, False),
            (" for unambiguous audit traces.", False, False),
        ],
        "link": [("RIC-349", False, False)],
    },
    {
        "cat": "", "cat_span": 0,
        "act": "Safety Landscape Tool — colleague self-service adoption + Dashboard updates",
        "desc": [
            ("1. Self-service intake app adopted by colleagues: ", False, False),
            ("10 submitted requests in August", True, False),
            (" (ZHCO, ZYGS, EJNH, YLV, SIWZ), incl. multi-target batches and MASH variants.", False, False),
            ("\n2. Dashboard: added ", False, False), ("Requester column", True, False),
            (" (fallback \"not know\") and \"Last 24 hours\" time filter.", False, False),
            ("\n3. Pre-check flows now flag archive-snapshot-only targets before filing.", False, False),
        ],
        "link": [("RIC-395", False, False)],
    },
    {
        "cat": "", "cat_span": 0,
        "act": "Imaging best practice — BP poster iteration",
        "desc": [
            ("Per leadership feedback the poster narrative was shifted from two case studies to ", False, False),
            ("general imaging best practice", True, False),
            (" (7-stage workflow + data flow), with assays demoted to proof-of-concept evidence; iterated to ", False, False),
            ("v49", True, False),
            (" through 6 review rounds.", False, False),
        ],
        "link": [("RIC-379", False, False)],
    },
    {
        "cat": "Targets", "cat_span": 2,
        "act": "GPR81 (HCAR1) individual target assessment — docking, reverse-SAR & safety",
        "desc": [
            ("1. Three-layer deliverable to requester UHYG: Vina docking | Boltz-2 | wet-lab benchmark, ", False, False),
            ("45/45 compounds", True, False),
            (" covered at every layer; full-chain data-integrity audit (20 checks, all pass).", False, False),
            ("\n2. New finding — ", False, False), ("binding-site disagreement between layers", True, False),
            (": Boltz-2 places 29/45 compounds in the orthosteric pocket contacting Arg71 (lactate carboxylate anchor); Vina on experimental structures disagrees \u2192 report section 7.", False, False),
            ("\n3. AlphaFold HCAR2/HCAR3 monomer models reproduced via BioLib and validated vs AlphaFold DB v6.", False, False),
            ("\n4. EN one-pager + readout email incl. wet-lab benchmarking panel ask; package synced to shared drive.", False, False),
        ],
        "link": [("RIC-396", False, False)],
    },
    {
        "cat": "", "cat_span": 0,
        "act": "Target safety pipeline — applications #3\u2013#8 logged & follow-ups",
        "desc": [
            ("1. CCND2 activation GREEN\u2192RED ", False, False), ("v1.3 offline rescore", True, False),
            (" (UHYG T2D batch, 42\u00d72 directions): keyword-library evolution flipped one Tier 1 COSMIC oncogene call.", False, False),
            ("\n2. CTRP9/FAM3D back-filled into registry + Dashboard (pre-request-form era run made visible).", False, False),
            ("\n3. Application #8: ", False, False), ("EJNH 3'aQTL panel, 19-target inhibition batch", True, False),
            (" under v1.4.3 scoring; pre-check flagged 17/19 as archive-snapshot-only before running.", False, False),
        ],
        "link": [("RIC-391", False, False)],
    },
    {
        "cat": "Assay", "cat_span": 3,
        "act": "IGWU myotube — factorial DOE + IGF/Insulin dose-response batches",
        "desc": [
            ("1. IGWU_20260730_1 Hypertrophy DOE: treatment \u00d7 arm joint-condition readout shipped in production decoupled v3.5 pipeline (", False, False),
            ("5,760 TIFFs", True, False),
            (", 240 wells); MSTN/Bima/IGF factors reported jointly instead of collapsed arms.", False, False),
            ("\n2. IGWU_20260811 ", False, False), ("4-plate onboard", True, False),
            (": IGF1/IGF2/Insulin dose series (Low/Mid) \u00d7 Atrophy/Hypertrophy arms; combined factorial report delivered.", False, False),
            ("\n3. On team follow-up ask, added MYH intensity readouts (SSMD) as report Section D + handed over raw readouts for their downstream analysis.", False, False),
        ],
        "link": [("RIC-394", False, False)],
    },
    {
        "cat": "", "cat_span": 0,
        "act": "Seahorse assay imaging pipeline — three development rounds",
        "desc": [
            ("1. Round 1: per-well Otsu masks produced holes from dark adipocyte clumps (nuclei signal inside holes ", False, False),
            ("5.6\u00d7", True, False),
            (" surrounding) \u2192 replaced with fixed-geometry mask templates.", False, False),
            ("\n2. Round 2: QC observation \"boundary hugs rim on some wells only\" triggered per-well geometry audit \u2014 fixed-radius template systematically wrong \u2192 switched to per-well radius templates.", False, False),
            ("\n3. Standalone seahorse_imaging project spun out.", False, False),
        ],
        "link": [("RIC-375", False, False)],
    },
    {
        "cat": "", "cat_span": 0,
        "act": "SkM Atrophy BP closeout + HepG2 ATP 4th modality",
        "desc": [
            ("1. SkM Atrophy methodology BP ", False, False), ("v2.1 summary", True, False),
            (": core-algorithm + code-map sections for engineering handover, metadata fields closed with assay-team input; code repository migrated to Novo GHE. (RIC-261, Review closeout)", False, False),
            ("\n2. HepG2 model characterization extended with ", False, False),
            ("ATP safety counter-screen as 4th modality", True, False),
            (" alongside TMRM/DRUG-seq/DINOv2 (Campaigns 31/24/22, NTC-normalised); one-pager final version + supplement delivered. (RIC-381)", False, False),
        ],
        "link": [("RIC-261", False, False), ("RIC-381", False, False)],
    },
]

VMERGE_SPANS = [
    (1, 3),
    (4, 2),
    (6, 3),
]

OUT = "reports/ric-monthly/2026-08/monthly_report_2026_08.pptx"

# ── Build ──
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_layout)

ML = Inches(0.25); MR = Inches(0.25); MT = Inches(0.15)
W = prs.slide_width; H = prs.slide_height

HDR_H = Inches(0.42)
hdr = slide.shapes.add_shape(1, ML, MT, W - ML - MR, HDR_H)
hdr.fill.solid(); hdr.fill.fore_color.rgb = NN_BLUE; hdr.line.fill.background()
tf = hdr.text_frame; tf.word_wrap = False
tf.margin_left = Inches(0.08); tf.margin_top = Inches(0.04)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
add_run(p, "Summary of activities", bold=True, size_pt=SZ_TITLE_BAR, color=WHITE)
add_run(p, "     \u2502     QYJI     \u2502     Page 1/1     \u2502     2026/08/28     \u2502     ", bold=False, size_pt=9, color=NN_BLUE_LITE)
add_run(p, "Novo Nordisk\u00ae", bold=True, size_pt=10, color=WHITE)

TBL_T = MT + HDR_H + Inches(0.06)
TBL_H = H - TBL_T - Inches(0.1)
TBL_W = W - ML - MR
N_ROWS = 1 + len(ROWS); N_COLS = 4

tbl_shape = slide.shapes.add_table(N_ROWS, N_COLS, ML, TBL_T, TBL_W, TBL_H)
tbl = tbl_shape.table

col_w = [int(TBL_W * 0.08), int(TBL_W * 0.15), int(TBL_W * 0.62), int(TBL_W * 0.15)]
col_w[-1] = TBL_W - sum(col_w[:-1])
for ci, w in enumerate(col_w):
    tbl.columns[ci].width = w

HDR_ROW_H  = Inches(0.28)
DATA_ROW_H = int((TBL_H - HDR_ROW_H) / max(len(ROWS), 1))
tbl.rows[0].height = HDR_ROW_H
for ri in range(1, N_ROWS):
    tbl.rows[ri].height = DATA_ROW_H

for ci, hh in enumerate(["Category", "List of activities", "Brief description", "Relevant links/files"]):
    cell = tbl.cell(0, ci)
    set_cell_bg(cell, NN_BLUE)
    fill_cell(cell, [(hh, True, False)], header=True)

for ri, row in enumerate(ROWS):
    dr = ri + 1
    bg = GREY_BG if dr % 2 == 0 else WHITE
    cat_cell = tbl.cell(dr, 0)
    set_cell_bg(cat_cell, NN_BLUE_LITE)
    fill_cell(cat_cell, [(row["cat"], True, False)] if row["cat"] else [("", False, False)], category=bool(row["cat"]))
    act_cell = tbl.cell(dr, 1)
    set_cell_bg(act_cell, bg); fill_cell(act_cell, [(row["act"], True, False)])
    desc_cell = tbl.cell(dr, 2)
    set_cell_bg(desc_cell, bg); fill_cell(desc_cell, row["desc"])
    link_cell = tbl.cell(dr, 3)
    set_cell_bg(link_cell, bg); fill_cell(link_cell, row["link"], link=True)

for (start, span) in VMERGE_SPANS:
    vmerge_restart(tbl.cell(start, 0))
    for i in range(1, span):
        vmerge_cont(tbl.cell(start + i, 0))

prs.save(OUT)
print("Saved:", OUT)
