"""
RIC Monthly Report — July 2026 (2026-06-26 -> 2026-07-27)
Built from templates/build_monthly_report.py (ric-monthly-report skill, confirmed 2026-06-25 fill_cell semantics).
Run: conda run -n ppt_env python scripts/ric-monthly/build_ric_monthly_2026_07.py
(from /das/user/QYJI/pulse repo root)
"""

import re

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

JIRA_BASE = "https://jira.novonordisk.com/browse/"
RIC_RE = re.compile(r"RIC-\d+")

# ── Colours ──
NN_BLUE      = RGBColor(0x00, 0x38, 0x65)
NN_BLUE_LITE = RGBColor(0xCC, 0xD9, 0xE8)
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GREY    = RGBColor(0x26, 0x26, 0x26)
GREY_BG      = RGBColor(0xF2, 0xF2, 0xF2)

# ── Font sizes (user-confirmed 2026-05-28, do not shrink) ──
SZ_HEADER    = 11
SZ_BODY      = 9.5
SZ_LINK      = 8.5
SZ_TITLE_BAR = 14
LINE_SPACING = 12

def set_cell_bg(cell, rgb):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for tag in ['a:solidFill', 'a:gradFill', 'a:noFill']:
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    solidFill = etree.SubElement(tcPr, qn('a:solidFill'))
    srgb = etree.SubElement(solidFill, qn('a:srgbClr'))
    srgb.set('val', f'{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}')

def add_run(para, text, bold=False, size_pt=SZ_BODY, color=DARK_GREY, italic=False, hyperlink=None):
    run = para.add_run()
    run.text = text
    run.font.bold = bold
    run.font.italic = italic
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.name = 'Calibri'
    if hyperlink:
        run.hyperlink.address = hyperlink
    return run

def _style_para(para):
    para.space_before = Pt(0.5)
    para.space_after  = Pt(0.5)
    para.line_spacing = Pt(LINE_SPACING)

def fill_cell(cell, lines, header=False, category=False, link=False):
    tf = cell.text_frame
    tf.word_wrap = True
    tf.margin_left   = Inches(0.04)
    tf.margin_right  = Inches(0.04)
    tf.margin_top    = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    c = WHITE if header else (NN_BLUE if category else DARK_GREY)
    sz = SZ_HEADER if header else (SZ_BODY if not link else SZ_LINK)

    para = tf.paragraphs[0]
    _style_para(para)
    para_used = False

    for (text, bold, italic) in lines:
        if link:
            if para_used:
                para = tf.add_paragraph(); _style_para(para)
            m = RIC_RE.search(text)
            href = JIRA_BASE + m.group(0) if m else None
            link_color = RGBColor(0x00, 0x56, 0x9E) if href else c
            add_run(para, text, bold=bold, italic=italic, size_pt=sz, color=link_color, hyperlink=href)
            para_used = True
            continue

        segments = text.split('\n')
        for si, seg in enumerate(segments):
            if si > 0:
                para = tf.add_paragraph(); _style_para(para)
                para_used = False
            if seg == '' and not para_used:
                continue
            add_run(para, seg, bold=bold, italic=italic, size_pt=sz, color=c)
            para_used = True

def vmerge_restart(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for el in tcPr.findall(qn('a:vMerge')):
        tcPr.remove(el)
    vm = etree.SubElement(tcPr, qn('a:vMerge'))
    vm.set('val', 'restart')

def vmerge_cont(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for el in tcPr.findall(qn('a:vMerge')):
        tcPr.remove(el)
    etree.SubElement(tcPr, qn('a:vMerge'))

# ── Table data ──
ROWS = [
    {
        "cat": "Capability", "cat_span": 3,
        "act": "Safety Landscape Tool — colleague self-service intake app",
        "desc": [
            ("1. New ", False, False), ("Dashboard", True, False),
            (": merges static archive index (", False, False), ("2,739 targets", True, False),
            (", 2026-07-06 snapshot) with live cron-run registries — colleagues self-check \"has this target been assessed?\" before filing a request.", False, False),
            ("\n2. ", False, False), ("Fixed Dashboard render bug", True, False),
            (" (jsonlite::validate() masking shiny::validate()).", False, False),
            ("\n3. ", False, False), ("Field simplification", True, False),
            (": retired High-precision/High-recall, kept only Standard (0.7082 binary accuracy); removed Assessment-type/Delivery-format dropdowns; merged name field into single 2-5 letter initial.", False, False),
            ("\n4. Live on Posit Connect, ", False, False), ("v1.2+35", True, False),
            (", with one-at-a-time entry + CSV bulk-add.", False, False),
        ],
        "link": [("RIC-395", False, False)],
    },
    {
        "cat": "", "cat_span": 0,
        "act": "Safety Landscape Tool — conservative RED operating point + science deep-dive + upstream alignment",
        "desc": [
            ("1. ", False, False), ("Default scoring switched to conservative RED call", True, False),
            (" (score>=75 AND max_weight>=60); reproduced NCBN-305 benchmark: binary accuracy ", False, False),
            ("0.7082", True, False), (", RED precision 0.6814, RED recall 0.5923.", False, False),
            ("\n2. ", False, False), ("34-slide science-meeting deep dive", True, False),
            (" (EN+CN): pipeline architecture, 8 data sources, keyword tuning history (Macro F1 0.381 -> 0.498), ~70% accuracy ceiling confirmed via 4 independent CV methods.", False, False),
            ("\n3. ", False, False), ("Benchmarked against Novo's in-house GHE safety-landscape-tool", True, False),
            (" (v0.91): 32/35 unit tests pass, 3 environment blockers identified, then ported keyword-scoring + dual-gate + NCBN-305 regression suite upstream.", False, False),
            ("\n4. request_form ", False, False), ("v1.2", True, False),
            (" live on Posit Connect: two-step submit with CSV/CLI preview, required direction field, third High-precision operating point added.", False, False),
        ],
        "link": [("RIC-349", False, False), ("RIC-364", False, False)],
    },
    {
        "cat": "", "cat_span": 0,
        "act": "HepG2 EE — cross-modal model characterization",
        "desc": [
            ("Extended cross-modal MoA story into a ", False, False), ("model characterization", True, False),
            (" package: DRUG-seq = target-perturbation coordinate system, TMRM/MitoTracker = mitochondrial functional readout, DINOv2 C24/C1 = image-derived state representation, pathway score = interpretable summary layer (not a replacement for full signature). Delivered pathway-vs-phenotype consistency analysis, tox-prediction benchmark scripts, and a one-pager.", False, False),
        ],
        "link": [("RIC-381", False, False)],
    },
    {
        "cat": "Targets", "cat_span": 2,
        "act": "Target safety assessment applications — SAM68 / HQGH 5-target batch",
        "desc": [
            ("1. ", False, False), ("SAM68 (KHDRBS1)", True, False),
            (" (req. SIWZ, T2D): both directions YELLOW (activation 67 / inhibition 77) — oncogenic splicing risk vs fertility/dyslipidemia signals; neither crosses RED.", False, False),
            ("\n2. ", False, False), ("HQGH batch", True, False),
            (" (5 targets, inhibition-only): BAIAP3 RED(82), KCNK16 YELLOW(52), P2RY1 RED(75), PKD1 RED(124, human ADPKD gene), TFRC RED(132, pLI 0.98 LoF-intolerant).", False, False),
        ],
        "link": [("RIC-391", False, False)],
    },
    {
        "cat": "", "cat_span": 0,
        "act": "GHSR inverse-agonist docking campaign — virtual screening + Boltz-2 cross-validation",
        "desc": [
            ("1. Dual-receptor protocol: primary dock vs inactive-state 7F83, counter-screen vs active-state 8JSR, ", False, False),
            ("7,931 ChEMBL", True, False), (" drug/lead-like compounds, 7,930/7,931 completed.", False, False),
            ("\n2. Classification: ", False, False), ("3,325 strong", True, False),
            (" inverse-agonist candidates (delta<-1.0), 2,268 moderate, 1,021 pan-state binders.", False, False),
            ("\n3. ", False, False), ("Top 50 cross-validated with Boltz-2", True, False),
            (" affinity prediction (50/50 done); top candidates AMINOQUINURIDE, MITOQUIDONE, ULECACICLIB.", False, False),
            ("\n4. Follow-up top-20 ChEMBL cross-check + liability screen scripts added.", False, False),
        ],
        "link": [("RIC-392", False, False)],
    },
    {
        "cat": "Assay", "cat_span": 1,
        "act": "IGWU myotube atrophy — batch application log (BP v2.0.0 on RIC-261 in Review closeout)",
        "desc": [
            ("1. ", False, False), ("IGWU_20260602_1 relabel + full rerun", True, False),
            (": biology team corrected 20-well plate layout (positive-control fix); only cell_analysis+readout rerun (CellProfiler/skeleton untouched) — surfaced SB431542 as a new strong hit (FiberWidth SSMD +2.23).", False, False),
            ("\n2. Onboarded new batch series ", False, False), ("IGWU_20260712", True, False),
            (" (Mid/Low/High dose) into the registry.", False, False),
            ("\n3. Ongoing cross-batch dashboard maintenance (dynamic STATS_CSV picker, renamed to multi_batch_screening_dashboard).", False, False),
        ],
        "link": [("RIC-394", False, False), ("RIC-261", False, False)],
    },
]

VMERGE_SPANS = [
    (1, 3),
    (4, 2),
    (6, 1),
]

OUT = "reports/ric-monthly/2026-07/monthly_report_2026_07.pptx"

# ── Build ──
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_layout)

ML = Inches(0.25); MR = Inches(0.25); MT = Inches(0.15)
W = prs.slide_width; H = prs.slide_height

HDR_H = Inches(0.42)
hdr = slide.shapes.add_shape(1, ML, MT, W - ML - MR, HDR_H)
hdr.fill.solid(); hdr.fill.fore_color.rgb = NN_BLUE; hdr.line.fill.background()
tf = hdr.text_frame; tf.word_wrap = False
tf.margin_left = Inches(0.08); tf.margin_top = Inches(0.04)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
add_run(p, "Summary of activities", bold=True, size_pt=SZ_TITLE_BAR, color=WHITE)
add_run(p, "     \u2502     QYJI     \u2502     Page 1/1     \u2502     2026/07/31     \u2502     ", bold=False, size_pt=9, color=NN_BLUE_LITE)
add_run(p, "Novo Nordisk\u00ae", bold=True, size_pt=10, color=WHITE)

TBL_T = MT + HDR_H + Inches(0.06)
TBL_H = H - TBL_T - Inches(0.1)
TBL_W = W - ML - MR
N_ROWS = 1 + len(ROWS); N_COLS = 4

tbl_shape = slide.shapes.add_table(N_ROWS, N_COLS, ML, TBL_T, TBL_W, TBL_H)
tbl = tbl_shape.table

col_w = [int(TBL_W * 0.08), int(TBL_W * 0.15), int(TBL_W * 0.62), int(TBL_W * 0.15)]
col_w[-1] = TBL_W - sum(col_w[:-1])
for ci, w in enumerate(col_w):
    tbl.columns[ci].width = w

HDR_ROW_H  = Inches(0.28)
DATA_ROW_H = int((TBL_H - HDR_ROW_H) / max(len(ROWS), 1))
tbl.rows[0].height = HDR_ROW_H
for ri in range(1, N_ROWS):
    tbl.rows[ri].height = DATA_ROW_H

for ci, hh in enumerate(["Category", "List of activities", "Brief description", "Relevant links/files"]):
    cell = tbl.cell(0, ci)
    set_cell_bg(cell, NN_BLUE)
    fill_cell(cell, [(hh, True, False)], header=True)

for ri, row in enumerate(ROWS):
    dr = ri + 1
    bg = GREY_BG if dr % 2 == 0 else WHITE
    cat_cell = tbl.cell(dr, 0)
    set_cell_bg(cat_cell, NN_BLUE_LITE)
    fill_cell(cat_cell, [(row["cat"], True, False)] if row["cat"] else [("", False, False)], category=bool(row["cat"]))
    act_cell = tbl.cell(dr, 1)
    set_cell_bg(act_cell, bg); fill_cell(act_cell, [(row["act"], True, False)])
    desc_cell = tbl.cell(dr, 2)
    set_cell_bg(desc_cell, bg); fill_cell(desc_cell, row["desc"])
    link_cell = tbl.cell(dr, 3)
    set_cell_bg(link_cell, bg); fill_cell(link_cell, row["link"], link=True)

for (start, span) in VMERGE_SPANS:
    vmerge_restart(tbl.cell(start, 0))
    for i in range(1, span):
        vmerge_cont(tbl.cell(start + i, 0))

prs.save(OUT)
print("Saved:", OUT)
